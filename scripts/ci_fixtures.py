# -*- coding: utf-8 -*-
"""CI 夹具：现场生成可转换的最小文件，并校验冻结包自测报告。

为什么需要它：发布分支只带运行所需的文件，不含 `tests/fixtures/`，而 CI 需要一组
确定性的输入来回答「冻结包还能不能转换」。这里全部用 **运行期已有依赖 + 标准库**
现场生成，不引入任何新依赖：

    xlsx -> openpyxl     pptx -> python-pptx     pdf -> Pillow
    docx -> zipfile      html/csv/json/xml/txt -> 标准库

用法::

    python scripts/ci_fixtures.py generate --out .ci-fixtures
    python scripts/ci_fixtures.py check ci-selftest.json --dir .ci-fixtures

@author 康小汪【kxw】
@date 2026-10-14
"""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

# 期望类型：text = 必须转出非空 Markdown；empty = 允许「成功但 0 字符」
EXPECT_TEXT = "text"
EXPECT_EMPTY = "empty"

# 标记只用字母数字：markitdown 会把下划线转义成 \_，含下划线会平白多一层坑
FIXTURES: dict[str, tuple[str, str]] = {
    "ci.docx": (EXPECT_TEXT, "CIGENDDOCX"),
    "ci.xlsx": (EXPECT_TEXT, "CIGENDXLSX"),
    "ci.pptx": (EXPECT_TEXT, "CIGENDPPTX"),
    "ci.html": (EXPECT_TEXT, "CIGENDHTML"),
    "ci.csv": (EXPECT_TEXT, "CIGENDCSV"),
    "ci.json": (EXPECT_TEXT, "CIGENDJSON"),
    "ci.xml": (EXPECT_TEXT, "CIGENDXML"),
    "ci.txt": (EXPECT_TEXT, "CIGENDTXT"),
    # Pillow 生成的 PDF 只有图像、没有文字层：这正是「转换成功但 0 字符」的真实路径，
    # 应用层据此给用户警告，CI 顺手把它钉住
    "ci.pdf": (EXPECT_EMPTY, ""),
}

TEXT_TEMPLATES: dict[str, str] = {
    "ci.html": "<html><body><h1>{marker}</h1><p>CI fixture</p></body></html>",
    "ci.csv": "name,value\n{marker},1\n",
    "ci.json": '{"{marker}": [1, 2, 3]}',
    "ci.xml": (
        '<?xml version="1.0" encoding="utf-8"?>'
        "<rss><channel><title>{marker}</title></channel></rss>"
    ),
    "ci.txt": "{marker}\n",
}

DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\
<Default Extension="xml" ContentType="application/xml"/>\
<Override PartName="/word/document.xml" \
ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>\
</Types>"""

DOCX_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\
<Relationship Id="rId1" \
Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" \
Target="word/document.xml"/>\
</Relationships>"""

DOCX_DOCUMENT = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">\
<w:body>\
<w:p><w:r><w:t>CI fixture</w:t></w:r></w:p>\
<w:p><w:r><w:t>{marker}</w:t></w:r></w:p>\
</w:body></w:document>"""


def write_docx(path: Path, marker: str) -> None:
    """手工拼最小 docx：mammoth 只认 [Content_Types].xml + _rels/.rels + word/document.xml。

    Args:
        path: 输出文件路径。
        marker: 写进正文的标记文本。
    """
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", DOCX_CONTENT_TYPES)
        archive.writestr("_rels/.rels", DOCX_RELS)
        archive.writestr("word/document.xml", DOCX_DOCUMENT.format(marker=marker))


def write_xlsx(path: Path, marker: str) -> None:
    """用 openpyxl 生成两行两列的 xlsx（依赖已是运行期依赖，不新增）。

    Args:
        path: 输出文件路径。
        marker: 写进单元格的标记文本。
    """
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "CI"
    sheet["A1"], sheet["B1"] = "key", "value"
    sheet["A2"], sheet["B2"] = marker, 42
    book.save(path)


def write_pptx(path: Path, marker: str) -> None:
    """用 python-pptx 生成一张带文本框的幻灯片。

    Args:
        path: 输出文件路径。
        marker: 写进文本框的标记文本。
    """
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    textbox.text_frame.text = marker
    deck.save(path)


def write_pdf(path: Path) -> None:
    """用 Pillow 生成一页「只有图像、没有文字层」的 PDF。

    Args:
        path: 输出文件路径。
    """
    from PIL import Image

    Image.new("RGB", (400, 200), "white").save(path, "PDF")


def write_text_fixture(path: Path, marker: str) -> None:
    """按后缀套模板写出纯文本类夹具。

    Args:
        path: 输出文件路径。
        marker: 写进内容的标记文本。
    """
    template = TEXT_TEMPLATES[path.name]
    path.write_text(template.replace("{marker}", marker), encoding="utf-8")


def generate(out_dir: Path) -> list[Path]:
    """生成全部夹具（重复调用幂等），返回文件列表。

    Args:
        out_dir: 输出目录，不存在时创建。

    Returns:
        list[Path]: 生成的文件路径，按文件名排序。
    """
    out_dir.mkdir(parents=True, exist_ok=True)
    for name, (_expect, marker) in FIXTURES.items():
        path = out_dir / name
        if name == "ci.docx":
            write_docx(path, marker)
        elif name == "ci.xlsx":
            write_xlsx(path, marker)
        elif name == "ci.pptx":
            write_pptx(path, marker)
        elif name == "ci.pdf":
            write_pdf(path)
        else:
            write_text_fixture(path, marker)
    return sorted(out_dir.iterdir())


def check_fixture_dir(fixture_dir: Path) -> list[str]:
    """校验夹具目录里就是预期的那批文件（缺文件、多文件都算问题）。

    Args:
        fixture_dir: 夹具目录。

    Returns:
        list[str]: 问题描述列表，空列表表示通过。
    """
    expected = set(FIXTURES)
    actual = {path.name for path in fixture_dir.iterdir() if path.is_file()}
    problems = [f"{name}: 夹具目录里缺文件" for name in sorted(expected - actual)]
    problems += [f"{name}: 夹具目录里多出未预期的文件" for name in sorted(actual - expected)]
    return problems


def _check_row(name: str, expect: str, row: dict[str, object]) -> list[str]:
    """校验报告里单个文件那一行。

    Args:
        name: 夹具文件名。
        expect: 期望类型（EXPECT_TEXT / EXPECT_EMPTY）。
        row: 自测报告里的一行。

    Returns:
        list[str]: 问题描述列表。
    """
    if not row.get("ok"):
        return [f"{name}: 转换失败，error={row.get('error') or '(空)'}"]
    chars = int(row.get("chars") or 0)
    if expect == EXPECT_EMPTY:
        if chars != 0:
            return [f"{name}: 期望 0 字符（图片型 PDF 无文字层），实际 {chars}"]
        return []
    if chars <= 0:
        return [f"{name}: 转换成功但结果是空的（0 字符）"]
    return []


def check_report(report_path: Path) -> list[str]:
    """校验冻结包自测报告：整体 ok + 每个夹具都符合预期类型。

    Args:
        report_path: `--selftest --json` 写出的报告路径。

    Returns:
        list[str]: 问题描述列表，空列表表示通过。
    """
    report = json.loads(report_path.read_text(encoding="utf-8"))
    rows = {Path(str(row.get("file", ""))).name: row for row in report.get("results", [])}
    problems = [] if report.get("ok") else ["报告整体 ok=False"]
    for name, (expect, _marker) in FIXTURES.items():
        row = rows.get(name)
        if row is None:
            problems.append(f"{name}: 报告里没有这一项")
            continue
        problems.extend(_check_row(name, expect, row))
    return problems


def _build_parser() -> argparse.ArgumentParser:
    """建命令行解析器。

    Returns:
        argparse.ArgumentParser: 带 generate / check 两个子命令的解析器。
    """
    parser = argparse.ArgumentParser(description="CI 夹具生成与校验")
    subparsers = parser.add_subparsers(dest="command", required=True)

    generate_parser = subparsers.add_parser("generate", help="生成夹具")
    generate_parser.add_argument("--out", default=".ci-fixtures", help="输出目录")

    check_parser = subparsers.add_parser("check", help="校验自测报告与夹具目录")
    check_parser.add_argument("report", help="自测报告 JSON 路径")
    check_parser.add_argument("--dir", default=".ci-fixtures", help="夹具目录")
    return parser


def main(argv: list[str] | None = None) -> int:
    """入口：generate 生成夹具，check 校验报告。

    Args:
        argv: 命令行参数，None 时取 sys.argv。

    Returns:
        int: 0 表示通过，1 表示有问题。
    """
    args = _build_parser().parse_args(argv)
    if args.command == "generate":
        paths = generate(Path(args.out))
        print(f"已生成 {len(paths)} 个夹具 -> {args.out}")
        return 0

    problems = check_fixture_dir(Path(args.dir)) + check_report(Path(args.report))
    for problem in problems:
        print(f"[FAIL] {problem}")
    print("夹具校验：全部通过" if not problems else f"夹具校验：{len(problems)} 项不合格")
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main())
